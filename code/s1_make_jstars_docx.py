#!/usr/bin/env python
# -*- coding: utf-8 -*-
# --- portable roots injected by the packaging script ---
import os as _os
PROJ_ROOT = _os.environ.get("S2DS_AUDIT_ROOT", _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
OUT_DIR   = _os.environ.get("S2DS_AUDIT_OUT", _os.path.join(PROJ_ROOT, "results"))
DATA_ROOT = _os.environ.get("S2DS_DATA_ROOT", _os.path.join(PROJ_ROOT, "data"))
# --- end injected ---

"""
P0-S1: render the JSTARS manuscript Markdown into a Word .docx for review.

Input : <ROOT>/P0S1_JSTARS_v1.md
Output: <ROOT>/P0S1_JSTARS_v1.docx

Single column and 1 inch margins on purpose: this copy is for reading and markup,
not for submission. The IEEEtran .tex next to it is the submission layout.
Fields the authors still have to supply are shown in red so they are easy to find.
"""
import os, re, io, sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = _os.path.join(PROJ_ROOT, "papers")
SRC = os.path.join(ROOT, "P0S1_JSTARS_v1.md")
DST = os.path.join(ROOT, "P0S1_JSTARS_v1.docx")
FIGDIR = os.path.join(ROOT, "figures")

FIG_FILES = {1: "fig1_noise_vs_testsize.png", 2: "fig2_effect_vs_noise.png"}

# Section numbers are printed as Roman numerals so the headings agree with the
# body cross-references ("Section IV-D"), matching the IEEEtran .tex layout.
ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII",
         8: "VIII", 9: "IX"}

BODY_FONT = "Times New Roman"
CJK_FONT = "宋体"

TITLE_SIZE = Pt(18)
AUTHOR_SIZE = Pt(11)
BODY_SIZE = Pt(10)
ABS_SIZE = Pt(9)
TBL_SIZE = Pt(8)
CAP_SIZE = Pt(8)
REF_SIZE = Pt(9)

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[\[.+?\]\]|<sup>.+?</sup>)", re.S)


# --------------------------------------------------------------------- helpers
def style_run(run, size=None, bold=None, italic=None, color=None, mono=False):
    name = "Courier New" if mono else BODY_FONT
    run.font.name = name
    run.font.size = size if size is not None else BODY_SIZE
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), CJK_FONT)
    return run


def clean(t):
    return t.replace("\\|", "|").replace("\\_", "_")


def add_runs(par, text, size=BODY_SIZE, bold=False, italic=False):
    """Add text to a paragraph, honouring **bold**, `code`, [[FILL]] and <sup>."""
    text = clean(text)
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            style_run(par.add_run(text[pos:m.start()]), size, bold, italic)
        tok = m.group(0)
        if tok.startswith("**"):
            style_run(par.add_run(clean(tok[2:-2])), size, True, italic)
        elif tok.startswith("`"):
            style_run(par.add_run(clean(tok[1:-1])), size, bold, italic, mono=True)
        elif tok.startswith("<sup>"):
            r = par.add_run(clean(tok[5:-6]))
            style_run(r, size, bold, italic)
            r.font.superscript = True
        else:  # [[FILL: ...]]
            r = par.add_run("[" + clean(tok[2:-2]) + "]")
            style_run(r, size, True, italic, RGBColor(0xC0, 0x00, 0x00))
        pos = m.end()
    if pos < len(text):
        style_run(par.add_run(text[pos:]), size, bold, italic)
    return par


def para(doc, text="", size=BODY_SIZE, bold=False, italic=False,
         align=None, space_after=6, space_before=0, first_line=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if first_line is not None:
        pf.first_line_indent = first_line
    if keep:
        pf.keep_with_next = True
    if text:
        add_runs(p, text, size, bold, italic)
    return p


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_page_numbers(section, label):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(label + "    "), Pt(8))
    run = p.add_run()
    style_run(run, Pt(8))
    for tag, attr, txt in (("w:fldChar", "w:fldCharType", "begin"),
                           ("w:instrText", "xml:space", "PAGE"),
                           ("w:fldChar", "w:fldCharType", "end")):
        el = OxmlElement(tag)
        el.set(qn(attr), txt)
        if tag == "w:instrText":
            el.text = " PAGE "
        run._r.append(el)


def split_row(raw):
    inner = raw.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|") and not inner.endswith("\\|"):
        inner = inner[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", inner)]


def add_table(doc, header, body):
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = TBL_SIZE if ncol >= 5 else Pt(8.5)

    hdr = t.rows[0]
    for j, txt in enumerate(header):
        cell = hdr.cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, txt, size, bold=True)
        shade(cell, "EDEDED")
    repeat_header(hdr)

    for row in body:
        cells = t.add_row().cells
        for j in range(ncol):
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            txt = row[j] if j < len(row) else ""
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, txt, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ------------------------------------------------------------------------ open
doc = Document()
st = doc.styles["Normal"]
st.font.name = BODY_FONT
st.font.size = BODY_SIZE
st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(1.0)
sec.left_margin = sec.right_margin = Inches(1.0)
add_page_numbers(sec, "P0-S1 JSTARS review copy")

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
n_tbl = n_fig = n_ref = n_bio = 0

while i < len(lines):
    s = lines[i].strip()

    if s.startswith("<!--"):
        while i < len(lines) and "-->" not in lines[i]:
            i += 1
        i += 1
        continue
    if not s:
        i += 1
        continue

    # ---------------------------------------------------------------- title
    if s.startswith("# "):
        p = para(doc, s[2:].strip(), TITLE_SIZE, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        i += 1
        continue

    # ------------------------------------------------------- author + thanks
    if re.match(r"^[A-Za-z].*<sup>1</sup>", s):
        para(doc, s, AUTHOR_SIZE, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        i += 1
        notes, orcid = [], ""
        while i < len(lines):
            t = lines[i].strip()
            if not t:
                i += 1
                continue
            if t.startswith("<sup>"):
                notes.append(re.sub(r"^<sup>.*?</sup>\s*", "", t).strip())
                i += 1
            elif t.startswith("ORCID:"):
                orcid = t
                i += 1
            else:
                break
        for t in notes:
            para(doc, t, Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        if orcid:
            para(doc, orcid, Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
        continue

    # ------------------------------------------------------------- abstract
    if s == "## Abstract":
        j = i + 1
        while j < len(lines) and not lines[j].strip():
            j += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        style_run(p.add_run("Abstract\u2014"), ABS_SIZE, bold=True, italic=True)
        add_runs(p, lines[j].strip(), ABS_SIZE)
        i = j + 1
        continue

    # --------------------------------------------------------- index terms
    if s.startswith("**Index Terms**"):
        kw = s.split("\u2014", 1)[-1]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        style_run(p.add_run("Index Terms\u2014"), ABS_SIZE, bold=True, italic=True)
        add_runs(p, kw.strip(), ABS_SIZE, italic=True)
        i += 1
        continue

    if s in ("---", "***", "___"):
        i += 1
        continue

    # ------------------------------------------------------ numbered section
    m = re.match(r"^##\s+(\d+)\.\s*(.+)$", s)
    if m:
        para(doc, f"{ROMAN[int(m.group(1))]}. {m.group(2).strip()}", Pt(11.5),
             bold=True, space_before=12, space_after=5, keep=True)
        i += 1
        continue

    # ------------------------------------------------ unnumbered / special
    if s.startswith("## "):
        name = s[3:].strip()

        if name == "References":
            para(doc, "References", Pt(11.5), bold=True, space_before=12,
                 space_after=5, keep=True)
            j = i + 1
            while j < len(lines):
                r = lines[j].strip()
                if not r:
                    j += 1
                    continue
                if r.startswith("##"):
                    break
                rm = re.match(r"^\[(\d+)\]\s*(.+)$", r)
                if rm:
                    n_ref += 1
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.space_after = Pt(3)
                    pf.line_spacing = 1.1
                    pf.left_indent = Inches(0.35)
                    pf.first_line_indent = Inches(-0.35)
                    style_run(p.add_run("[%s]  " % rm.group(1)), REF_SIZE)
                    add_runs(p, rm.group(2), REF_SIZE)
                j += 1
            i = j
            continue

        if name == "Author Biographies":
            para(doc, "Author Biographies", Pt(11.5), bold=True, space_before=14,
                 space_after=5, keep=True)
            j = i + 1
            while j < len(lines):
                r = lines[j].strip()
                if not r:
                    j += 1
                    continue
                if r.startswith("##"):
                    break
                nm = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", r)
                if nm and nm.group(1) in ("Zhengyang Chen", "Dawei Liu", "Zhiheng Liu"):
                    n_bio += 1
                    p = doc.add_paragraph()
                    pf = p.paragraph_format
                    pf.space_after = Pt(8)
                    pf.line_spacing = 1.15
                    pf.left_indent = Inches(0.55)
                    pf.first_line_indent = Inches(-0.55)
                    style_run(p.add_run(nm.group(1) + "  "), BODY_SIZE, bold=True)
                    add_runs(p, nm.group(2), BODY_SIZE)
                else:
                    para(doc, r, Pt(9), italic=True)
                j += 1
            i = j
            continue

        if name == "Figures and Tables":
            para(doc, name, Pt(11.5), bold=True, space_before=14, space_after=5, keep=True)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(lines[j])
                j += 1
            cells = [split_row(r) for r in rows]
            add_table(doc, cells[0], cells[2:])
            i = j
            continue

        para(doc, name, Pt(11.5), bold=True, space_before=12, space_after=5, keep=True)
        i += 1
        continue

    # ---------------------------------------------------------- subsection
    m = re.match(r"^###\s+([A-Z])\.\s*(.+)$", s)
    if m:
        para(doc, f"{m.group(1)}. {m.group(2).strip()}", Pt(10), bold=True,
             italic=True, space_before=9, space_after=4, keep=True)
        i += 1
        continue

    # -------------------------------------------------------- table caption
    m = re.match(r"^\*\*TABLE\s+([IVX]+)\.\s*(.+?)\*\*$", s)
    if m:
        num, cap = m.group(1), m.group(2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        style_run(p.add_run("TABLE %s" % num), CAP_SIZE, bold=True)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.keep_with_next = True
        add_runs(p2, cap, CAP_SIZE)
        j = i + 1
        while j < len(lines) and not lines[j].strip():
            j += 1
        rows = []
        while j < len(lines) and lines[j].strip().startswith("|"):
            rows.append(lines[j])
            j += 1
        cells = [split_row(r) for r in rows]
        add_table(doc, cells[0], cells[2:])
        n_tbl += 1
        i = j
        continue

    # ------------------------------------------------------- figure caption
    m = re.match(r"^\*\*Fig\.\s+(\d+)\.\s*(.+?)\*\*\s*(.*)$", s)
    if m:
        num = int(m.group(1))
        cap, rest = m.group(2), m.group(3).strip()
        fpath = os.path.join(FIGDIR, FIG_FILES.get(num, ""))
        if os.path.exists(fpath):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.add_run().add_picture(fpath, width=Inches(5.6))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        style_run(p.add_run("Fig. %d.  " % num), CAP_SIZE, bold=True)
        add_runs(p, cap, CAP_SIZE)
        n_fig += 1
        if rest:
            para(doc, rest, BODY_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
        i += 1
        continue

    # -------------------------------------------------------- numbered list
    if re.match(r"^\d{1,2}\.\s", s):
        while i < len(lines) and re.match(r"^\d{1,2}\.\s", lines[i].strip()):
            txt = re.sub(r"^\d{1,2}\.\s", "", lines[i].strip())
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_runs(p, txt, BODY_SIZE)
            i += 1
        continue

    # ------------------------------------------------------------- paragraph
    para(doc, s, BODY_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    i += 1

doc.save(DST)
print("wrote %s (%d bytes)" % (DST, os.path.getsize(DST)))
print("tables=%d figures=%d references=%d biographies=%d" % (n_tbl, n_fig, n_ref, n_bio))

# python-docx stamps its own default template onto every file it writes. Strip
# those traces so the review copy carries no tool metadata.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from s1_scrub_docx_metadata import scrub as _scrub
_scrub(DST)
print("metadata scrubbed: %s" % os.path.basename(DST))
